"""文档解析：把各类文件转成带页码/标题层级的文本块。"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl", ".yaml", ".yml",
    ".log", ".ini", ".conf", ".toml", ".sql", ".sh", ".bat", ".ps1",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".vue", ".java", ".go", ".rs", ".c", ".h",
    ".cpp", ".hpp", ".cs", ".php", ".rb", ".kt", ".swift", ".scala", ".r", ".m",
    ".xml", ".properties", ".gradle", ".dockerfile", ".env.example",
}
OFFICE_EXTS = {".docx", ".xlsx", ".pptx"}
HTML_EXTS = {".html", ".htm"}
SUPPORTED_EXTS = {".pdf"} | OFFICE_EXTS | HTML_EXTS | TEXT_EXTS


class ParseError(RuntimeError):
    pass


@dataclass
class Block:
    text: str
    page: int | None = None
    heading: str | None = None


@dataclass
class ParsedDoc:
    title: str
    blocks: list[Block] = field(default_factory=list)
    meta: dict = field(default_factory=dict)

    @property
    def chars(self) -> int:
        return sum(len(b.text) for b in self.blocks)


def clean_text(text: str) -> str:
    if not text:
        return ""
    text = text.replace("\ufeff", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def parse_file(path: Path) -> ParsedDoc:
    ext = path.suffix.lower()
    title = path.stem
    if ext == ".pdf":
        doc = _parse_pdf(path)
    elif ext == ".docx":
        doc = _parse_docx(path)
    elif ext == ".xlsx":
        doc = _parse_xlsx(path)
    elif ext == ".pptx":
        doc = _parse_pptx(path)
    elif ext in HTML_EXTS:
        doc = _parse_html(path.read_text(encoding="utf-8", errors="ignore"), title)
    elif ext in {".csv", ".tsv"}:
        doc = _parse_csv(path)
    elif ext in TEXT_EXTS:
        doc = _parse_text(path.read_text(encoding="utf-8", errors="ignore"), title)
    else:
        raise ParseError(f"暂不支持的文件类型：{ext or '无扩展名'}")
    doc.title = doc.title or title
    if doc.chars < 30:
        raise ParseError("解析出的文本过少，可能是扫描件/图片型文档，暂不支持 OCR")
    return doc


def _parse_pdf(path: Path) -> ParsedDoc:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"PDF 打开失败：{exc}") from exc
    blocks: list[Block] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = clean_text(page.extract_text() or "")
        except Exception:  # noqa: BLE001
            text = ""
        if text:
            blocks.append(Block(text=text, page=index))
    meta = {"pages": len(reader.pages)}
    title = ""
    try:
        info = reader.metadata or {}
        title = str(info.get("/Title") or "").strip()
    except Exception:  # noqa: BLE001
        title = ""
    return ParsedDoc(title=title or path.stem, blocks=blocks, meta=meta)


def _docx_blocks(container) -> list[Block]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    blocks: list[Block] = []
    stack: list[tuple[int, str]] = []

    def current_heading() -> str | None:
        return " > ".join(t for _, t in stack) if stack else None

    body = container.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, container)
            text = clean_text(para.text)
            if not text:
                continue
            style = (para.style.name or "") if para.style is not None else ""
            match = re.match(r"(?:Heading|标题)\s*(\d)", style)
            level = int(match.group(1)) if match else (1 if style.lower() in {"title", "标题"} else 0)
            if level:
                while stack and stack[-1][0] >= level:
                    stack.pop()
                stack.append((level, text))
                blocks.append(Block(text=text, heading=current_heading()))
            else:
                blocks.append(Block(text=text, heading=current_heading()))
        elif tag == "tbl":
            table = Table(child, container)
            rows = []
            for row in table.rows:
                cells = [clean_text(cell.text).replace("\n", " ") for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                blocks.append(Block(text="\n".join(rows), heading=current_heading()))
    return blocks


def _parse_docx(path: Path) -> ParsedDoc:
    import docx

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"Word 打开失败：{exc}") from exc
    return ParsedDoc(title=path.stem, blocks=_docx_blocks(document))


MAX_XLSX_ROWS = 3000
MAX_XLSX_COLS = 24

# csv 走结构化渲染，避免行内逗号污染分块边界


def _parse_xlsx(path: Path) -> ParsedDoc:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"Excel 打开失败：{exc}") from exc
    blocks: list[Block] = []
    for sheet in workbook.worksheets:
        lines: list[str] = []
        row_count = 0
        for row in sheet.iter_rows(values_only=True):
            values = ["" if v is None else str(v).strip() for v in row[:MAX_XLSX_COLS]]
            if any(values):
                lines.append(" | ".join(values).rstrip(" |"))
                row_count += 1
            if row_count >= MAX_XLSX_ROWS:
                lines.append(f"（超出 {MAX_XLSX_ROWS} 行，已截断）")
                break
        if lines:
            blocks.append(Block(text=clean_text("\n".join(lines)), heading=sheet.title))
    sheet_count = len(workbook.sheetnames)
    workbook.close()
    return ParsedDoc(title=path.stem, blocks=blocks, meta={"sheets": sheet_count})


def _parse_pptx(path: Path) -> ParsedDoc:
    from pptx import Presentation

    try:
        prs = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"PPT 打开失败：{exc}") from exc
    blocks: list[Block] = []
    for index, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        title = ""
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = clean_text(shape.text_frame.text)
            if not text:
                continue
            if not title and shape == slide.shapes.title:
                title = text.split("\n")[0][:80]
            parts.append(text)
        if parts:
            blocks.append(Block(text="\n".join(parts), page=index, heading=title or f"第 {index} 页"))
    return ParsedDoc(title=path.stem, blocks=blocks, meta={"slides": len(prs.slides)})


def _parse_text(raw: str, title: str) -> ParsedDoc:
    blocks: list[Block] = []
    stack: list[tuple[int, str]] = []
    buffer: list[str] = []
    heading: str | None = None

    def flush() -> None:
        nonlocal buffer
        text = clean_text("\n".join(buffer))
        if text:
            blocks.append(Block(text=text, heading=heading))
        buffer = []

    for line in raw.replace("\r\n", "\n").split("\n"):
        match = re.match(r"^(#{1,6})\s+(.+)$", line)
        if match:
            flush()
            level = len(match.group(1))
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, match.group(2).strip()))
            heading = " > ".join(t for _, t in stack)
            buffer.append(match.group(2).strip())
        else:
            buffer.append(line)
    flush()
    if not blocks:
        blocks = [Block(text=clean_text(raw))]
    return ParsedDoc(title=title, blocks=blocks)


def _parse_csv(path: Path) -> ParsedDoc:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    reader = csv.reader(io.StringIO(raw))
    lines = [" | ".join(cell.strip() for cell in row) for row in reader if any(c.strip() for c in row)]
    return ParsedDoc(title=path.stem, blocks=[Block(text=clean_text("\n".join(lines)))], meta={"rows": len(lines)})


def _parse_html(raw: str, title: str) -> ParsedDoc:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(raw, "lxml")
    page_title = ""
    if soup.title and soup.title.string:
        page_title = soup.title.string.strip()
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "form", "aside", "svg"]):
        tag.decompose()
    blocks: list[Block] = []
    heading = ""
    for node in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "pre", "td"]):
        text = clean_text(node.get_text(" ", strip=True))
        if not text or len(text) < 2:
            continue
        if node.name in {"h1", "h2", "h3", "h4"}:
            heading = text[:120]
            blocks.append(Block(text=text, heading=heading))
        else:
            blocks.append(Block(text=text, heading=heading or None))
    return ParsedDoc(title=page_title or title, blocks=blocks)


def fetch_url(url: str, timeout: float = 25.0) -> ParsedDoc:
    import httpx

    if not re.match(r"^https?://", url):
        raise ParseError("只支持 http/https 链接")
    headers = {"User-Agent": "Mozilla/5.0 (compatible; LocalRAG/0.1)"}
    try:
        with httpx.Client(follow_redirects=True, timeout=timeout, headers=headers) as client:
            resp = client.get(url)
            resp.raise_for_status()
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"抓取失败：{exc}") from exc
    ctype = resp.headers.get("content-type", "")
    if "pdf" in ctype:
        raise ParseError("该链接是 PDF，请下载后上传")
    if "html" not in ctype and "text" not in ctype:
        raise ParseError(f"不支持的内容类型：{ctype}")
    doc = _parse_html(resp.text, url)
    doc.meta["url"] = url
    return doc
